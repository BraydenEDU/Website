#!/usr/bin/env python3
"""Build the Annuit marketing site into static HTML.

    python build.py            # writes dist/  (deploy this folder)
    python build.py --serve    # builds, then serves dist/ at http://localhost:8000
    python build.py --preview  # writes preview/ for the shareable Claude preview
    python build.py --copy-doc # exports all copy to docs/annuit-site-copy.docx for review

Content lives in content/<lang>.yaml, settings in site.yaml, layout in templates/.
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import functools
import http.server
import posixpath
import shutil
import socketserver
import sys
from pathlib import Path

import yaml
from jinja2 import Environment, FileSystemLoader, StrictUndefined

ROOT = Path(__file__).resolve().parent

# page slug -> template file
PAGES = {
    "index": "home.html",
    "pricing": "pricing.html",
    "security": "security.html",
    "demo": "demo.html",
    "privacy": "legal.html",
    "terms": "legal.html",
}


def load_yaml(path: Path) -> dict:
    with path.open(encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def deep_merge(base: dict, override: dict) -> dict:
    """Override wins; missing keys fall back to base (English)."""
    out = copy.deepcopy(base)
    for key, value in (override or {}).items():
        if isinstance(value, dict) and isinstance(out.get(key), dict):
            out[key] = deep_merge(out[key], value)
        else:
            out[key] = value
    return out


def page_path(lang: str, page: str, default_lang: str) -> str:
    prefix = "" if lang == default_lang else f"{lang}/"
    return f"{prefix}{page}.html"


def make_helpers(current: str, default_lang: str, preview: bool):
    cur_dir = posixpath.dirname(current) or "."

    def url(page: str, lang: str, anchor: str = "") -> str:
        target = page_path(lang, page, default_lang)
        if preview and target == "index.html":
            # The preview host serves the English home page at the root.
            rel = posixpath.relpath(".", cur_dir)
            href = "./" if rel == "." else rel + "/"
        else:
            href = posixpath.relpath(target, cur_dir)
        return href + (f"#{anchor}" if anchor else "")

    def asset(name: str) -> str:
        return posixpath.relpath(f"assets/{name}", cur_dir)

    return url, asset


def fmt(template: str, **values) -> str:
    """Fill {placeholders} in a content string."""
    return template.format(**values)


def build(out_dir: Path, preview: bool = False) -> list[str]:
    site = load_yaml(ROOT / "site.yaml")
    default_lang = site["default_lang"]
    english = load_yaml(ROOT / "content" / f"{default_lang}.yaml")

    env = Environment(
        loader=FileSystemLoader(ROOT / "templates"),
        undefined=StrictUndefined,
        autoescape=True,
        trim_blocks=True,
        lstrip_blocks=True,
    )
    env.filters["fmt"] = lambda s, **kw: fmt(s, **kw)

    if out_dir.exists():
        shutil.rmtree(out_dir)
    out_dir.mkdir(parents=True)

    docs = site["demo_client"]["documents"]
    counts = {s: sum(1 for d in docs if d["status"] == s) for s in ("received", "wrong_year", "missing")}

    written = []
    for lang in site["languages"]:
        lang_file = ROOT / "content" / f"{lang}.yaml"
        t = english if lang == default_lang else deep_merge(english, load_yaml(lang_file))
        for page, template_name in PAGES.items():
            rel_path = page_path(lang, page, default_lang)
            url, asset = make_helpers(rel_path, default_lang, preview)
            html = env.get_template(template_name).render(
                s=site,
                t=t,
                lang=lang,
                page=page,
                url=url,
                asset=asset,
                counts=counts,
                preview=preview,
                fragment=preview and rel_path == "index.html",
                year=dt.date.today().year,
                built=dt.date.today().isoformat(),
            )
            target = out_dir / rel_path
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(html, encoding="utf-8")
            written.append(rel_path)

    # 404 page (English) at the root; most static hosts serve 404.html automatically.
    url, asset = make_helpers("404.html", default_lang, preview)
    (out_dir / "404.html").write_text(env.get_template("404.html").render(
        s=site, t=english, lang=default_lang, page="index", url=url, asset=asset, counts=counts,
        preview=preview, fragment=False, year=dt.date.today().year, built=dt.date.today().isoformat()), encoding="utf-8")
    written.append("404.html")

    if not preview:
        written += write_seo_files(out_dir, site)

    shutil.copytree(ROOT / "static", out_dir / "assets")
    written += [f"assets/{p.relative_to(ROOT / 'static').as_posix()}" for p in (ROOT / "static").rglob("*") if p.is_file()]
    return written


def public_url(site: dict, lang: str, page: str) -> str:
    prefix = "" if lang == site["default_lang"] else f"{lang}/"
    return f"https://{site['domain']}/{prefix}{'' if page == 'index' else page}"


def write_seo_files(out_dir: Path, site: dict) -> list[str]:
    """sitemap.xml with hreflang alternates, plus robots.txt. URLs are extensionless (vercel.json cleanUrls)."""
    today = dt.date.today().isoformat()
    rows = []
    for lang in site["languages"]:
        for page in PAGES:
            alts = "".join(
                f'\n    <xhtml:link rel="alternate" hreflang="{site["languages"][alt]["html_lang"]}" href="{public_url(site, alt, page)}"/>'
                for alt in site["languages"]
            )
            rows.append(f"  <url>\n    <loc>{public_url(site, lang, page)}</loc>\n    <lastmod>{today}</lastmod>{alts}\n  </url>")
    xml = ('<?xml version="1.0" encoding="UTF-8"?>\n'
           '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9" xmlns:xhtml="http://www.w3.org/1999/xhtml">\n'
           + "\n".join(rows) + "\n</urlset>\n")
    (out_dir / "sitemap.xml").write_text(xml, encoding="utf-8")
    (out_dir / "robots.txt").write_text(f"User-agent: *\nAllow: /\nSitemap: https://{site['domain']}/sitemap.xml\n", encoding="utf-8")
    return ["sitemap.xml", "robots.txt"]


def write_copy_doc(path: Path) -> None:
    """Export every content string to a Word doc Annuit can comment on (deliverable 2)."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    site = load_yaml(ROOT / "site.yaml")
    english = load_yaml(ROOT / "content" / f"{site['default_lang']}.yaml")
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Annuit website copy", 0)
    doc.add_paragraph(f"Generated {dt.date.today().isoformat()} from content/*.yaml. Comment directly in this doc; "
                      "approved changes go back into the YAML files. Grey text is the content key.")

    def walk(node, key_path, langs):
        if isinstance(node, dict):
            for k, v in node.items():
                walk(v, key_path + [k], langs)
        elif isinstance(node, list):
            for i, v in enumerate(node, 1):
                walk(v, key_path + [str(i)], langs)
        else:
            p = doc.add_paragraph()
            run = p.add_run(".".join(key_path) + "\n")
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            p.add_run(str(node))
            for code, data in langs.items():
                cur = data
                for k in key_path:
                    cur = cur[int(k) - 1] if isinstance(cur, list) else (cur or {}).get(k) if isinstance(cur, dict) else None
                    if cur is None:
                        break
                if cur is not None:
                    r = p.add_run(f"\n[{code.upper()}] {cur}")
                    r.font.color.rgb = RGBColor(0x55, 0x44, 0x3A)

    others = {c: load_yaml(ROOT / "content" / f"{c}.yaml") for c in site["languages"] if c != site["default_lang"]}
    for section, body in english.items():
        doc.add_heading(section.replace("_", " ").title(), 1)
        walk(body, [section], others)

    doc.add_heading("Site settings shown to visitors", 1)
    for label, value in [("Customers (trust line)", ", ".join(site["customers"])),
                         ("Testimonial", f'{site["testimonial"]["quote"]} | {site["testimonial"]["name"]}, {site["testimonial"]["title"]}')] + \
                        [(f"Follow-up message ({c})", m) for c, m in site["follow_up_messages"].items()]:
        p = doc.add_paragraph()
        p.add_run(label + ": ").bold = True
        p.add_run(value)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def serve(directory: Path, port: int = 8000) -> None:
    handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=str(directory))
    with socketserver.TCPServer(("", port), handler) as httpd:
        print(f"Serving {directory} at http://localhost:{port}  (Ctrl+C to stop)")
        httpd.serve_forever()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--preview", action="store_true", help="build the Claude preview variant into preview/")
    parser.add_argument("--serve", action="store_true", help="serve the build locally after building")
    parser.add_argument("--port", type=int, default=8000)
    parser.add_argument("--copy-doc", action="store_true", help="export all site copy to docs/annuit-site-copy.docx")
    args = parser.parse_args()

    if args.copy_doc:
        target = ROOT / "docs" / "annuit-site-copy.docx"
        write_copy_doc(target)
        print(f"Wrote {target.relative_to(ROOT)}")
        return 0

    out = ROOT / ("preview" if args.preview else "dist")
    files = build(out, preview=args.preview)
    print(f"Built {len(files)} files into {out.relative_to(ROOT)}/")
    if args.serve:
        serve(out, args.port)
    return 0


if __name__ == "__main__":
    sys.exit(main())
